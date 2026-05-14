from openpyxl import Workbook

class Excel:
    def __init__(self, x, y, estatisticas):
        self.x = x
        self.y = y
        self.estatisticas = estatisticas

    def exportar(self, nome_arquivo):

        # cria planilha
        wb = Workbook()

        ws = wb.active

        ws.title = "Dados"

        #cabecalho
        ws['A1'] = "X"
        ws['B1'] = "Y"

        # dados
        for i in range(len(self.y)):

            if self.x:
                ws.cell(row=i + 2, column=1, value=self.x[i])

            ws.cell(row=i + 2, column=2, value=self.y[i])

        # nova aba estatísticas
        ws2 = wb.create_sheet("Estatísticas")

        ws2["A1"] = "Estatística"

        ws2["B1"] = "Valor"

        ws2["A2"] = "Média"
        ws2["B2"] = self.estatisticas.media()

        ws2["A3"] = "Maior"
        ws2["B3"] = self.estatisticas.maior()

        ws2["A4"] = "Menor"
        ws2["B4"] = self.estatisticas.menor()

        ws2["A5"] = "Soma"
        ws2["B5"] = self.estatisticas.soma()

        ws2["A6"] = "Desvio Padrão"
        ws2["B6"] = self.estatisticas.desvio_padrao()

        # salva arquivo
        wb.save(f"{nome_arquivo}.xlsx")